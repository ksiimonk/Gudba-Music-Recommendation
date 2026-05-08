import docx

path = r'C:\dev\tasks\Gudba-Music-Recommendation\VKR_Gudba_Music_Recommendation.docx'
doc = docx.Document(path)

# Look at paragraphs around chapter headings
print("=== Sample paragraphs ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    # Print heading-like paragraphs and their neighbors
    if any(t.startswith(x) for x in ['1 ', '2 ', '3 ', '1.', '2.', '3.', 'Введ', 'Закл', 'Спис']):
        print(f'\n[{i}] HEADING: {t[:100]}')
        # Print next paragraph
        if i+1 < len(doc.paragraphs):
            n = doc.paragraphs[i+1].text.strip()[:80]
            print(f'  [{i+1}] NEXT: {n}')
        if i+2 < len(doc.paragraphs):
            n = doc.paragraphs[i+2].text.strip()[:80]
            print(f'  [{i+2}] NEXT2: {n}')

print("\n=== Content check (first 80 chars of notable paragraphs) ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if len(t) > 500:
        print(f'[{i}] {t[:80]}... ({len(t)} chars)')
